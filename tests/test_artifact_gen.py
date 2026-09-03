from __future__ import annotations

import json
import os
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import patch

from partner.office import artifact_gen as af

TABLE_MD = """| 城市 | 销量 | 占比 |
| --- | --- | --- |
| 北京 | 1,200 | 45% |
| 上海 | 800 | 30% |
"""

TWO_TABLES_MD = (
    TABLE_MD
    + "\n说明文字\n\n| 月份 | 收入 |\n| --- | --- |\n| 一月 | 10 |\n| 二月 | 20 |\n"
)


class TestParseTables(unittest.TestCase):
    def test_single_table(self):
        tables = af.parse_markdown_tables(TABLE_MD)
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0][0], ["城市", "销量", "占比"])
        self.assertEqual(tables[0][1], ["北京", "1,200", "45%"])

    def test_multiple_tables(self):
        tables = af.parse_markdown_tables(TWO_TABLES_MD)
        self.assertEqual(len(tables), 2)
        self.assertEqual(tables[1][0], ["月份", "收入"])

    def test_no_table(self):
        self.assertEqual(af.parse_markdown_tables("没有表格的文本"), [])

    def test_header_only_rejected(self):
        md = "| a | b |\n| --- | --- |\n"
        self.assertEqual(af.parse_markdown_tables(md), [])


class TestMaybeNumber(unittest.TestCase):
    def test_plain_number(self):
        self.assertEqual(af._maybe_number("42"), 42)
        self.assertEqual(af._maybe_number("3.5"), 3.5)

    def test_thousands_separator(self):
        self.assertEqual(af._maybe_number("1,200"), 1200)

    def test_percent_to_fraction(self):
        self.assertEqual(af._maybe_number("45%"), 0.45)

    def test_text_passthrough(self):
        self.assertEqual(af._maybe_number("北京"), "北京")


class TestBuildXlsx(unittest.TestCase):
    def test_build_and_readback(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch.dict(os.environ, {"FEISHU_PARTNER_ARTIFACT_DIR": tmp}):
                tables = af.parse_markdown_tables(TABLE_MD)
                path = af.build_xlsx(tables, filename="销量")
                self.assertTrue(path.exists())
                self.assertEqual(path.name, "销量.xlsx")
                from openpyxl import load_workbook

                wb = load_workbook(str(path))
                sheet = wb.active
                self.assertEqual(sheet.cell(1, 1).value, "城市")
                self.assertEqual(sheet.cell(2, 2).value, 1200)
                self.assertAlmostEqual(sheet.cell(2, 3).value, 0.45)

    def test_multi_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch.dict(os.environ, {"FEISHU_PARTNER_ARTIFACT_DIR": tmp}):
                tables = af.parse_markdown_tables(TWO_TABLES_MD)
                path = af.build_xlsx(tables)
                from openpyxl import load_workbook

                wb = load_workbook(str(path))
                self.assertEqual(len(wb.sheetnames), 2)


class TestMakeExcel(unittest.TestCase):
    def test_no_table_message(self):
        result = af.make_excel("随便一段文字")
        self.assertIn("没解析出", result)

    def test_generates_without_upload(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch.dict(os.environ, {"FEISHU_PARTNER_ARTIFACT_DIR": tmp}):
                result = af.make_excel(TABLE_MD)
        self.assertIn("已生成", result)
        self.assertIn("未上传", result)

    def test_at_file_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            md = Path(tmp) / "t.md"
            md.write_text(TABLE_MD, encoding="utf-8")
            with patch.dict(os.environ, {"FEISHU_PARTNER_ARTIFACT_DIR": tmp}):
                result = af.make_excel(f"@{md}")
        self.assertIn("已生成", result)

    def test_upload_calls_import(self):
        payload = {"ok": True, "data": {"result": {"url": "https://x.feishu.cn/sheets/abc"}}}
        with tempfile.TemporaryDirectory() as tmp:
            with patch.dict(os.environ, {"FEISHU_PARTNER_ARTIFACT_DIR": tmp}):
                with patch.object(af, "run_lark", return_value=payload) as mock_run:
                    result = af.make_excel(TABLE_MD, upload=True, name="销量表")
        self.assertIn("https://x.feishu.cn/sheets/abc", result)
        argv = mock_run.call_args[0][0]
        self.assertIn("--type", argv)
        self.assertIn("sheet", argv)


class TestLastAnswerStore(unittest.TestCase):
    def test_roundtrip(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Path(tmp) / "la.json"
            with patch.dict(os.environ, {"FEISHU_PARTNER_LAST_ANSWER": str(store)}):
                af.save_last_answer("oc_1", "本周销量如何", "本周销量 2000 台，环比增长 20%。")
                row = af.load_last_answer("oc_1")
        self.assertIsNotNone(row)
        self.assertIn("2000", row["answer"])
        self.assertEqual(row["query"], "本周销量如何")

    def test_short_answer_not_saved(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Path(tmp) / "la.json"
            with patch.dict(os.environ, {"FEISHU_PARTNER_LAST_ANSWER": str(store)}):
                af.save_last_answer("oc_1", "q", "好")
                self.assertIsNone(af.load_last_answer("oc_1"))

    def test_expired_answer_dropped(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Path(tmp) / "la.json"
            old = (datetime.now(af.CN_TZ) - timedelta(hours=3)).isoformat(timespec="seconds")
            store.write_text(
                json.dumps({"oc_1": {"ts": old, "query": "q", "answer": "x" * 50}}),
                encoding="utf-8",
            )
            with patch.dict(os.environ, {"FEISHU_PARTNER_LAST_ANSWER": str(store)}):
                self.assertIsNone(af.load_last_answer("oc_1"))

    def test_unknown_chat(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Path(tmp) / "la.json"
            with patch.dict(os.environ, {"FEISHU_PARTNER_LAST_ANSWER": str(store)}):
                self.assertIsNone(af.load_last_answer("oc_none"))


class TestDocTrigger(unittest.TestCase):
    def test_trigger_phrases(self):
        for phrase in ("转文档", "转成文档", "生成文档", "存成文档", "转文档！"):
            self.assertTrue(af.looks_like_doc_request(phrase), phrase)

    def test_non_trigger(self):
        self.assertFalse(af.looks_like_doc_request("帮我把周报写成文档然后发给所有人审阅一下"))
        self.assertFalse(af.looks_like_doc_request("今天有什么任务"))


class TestAnswerToDoc(unittest.TestCase):
    def test_no_answer_cached(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Path(tmp) / "la.json"
            with patch.dict(os.environ, {"FEISHU_PARTNER_LAST_ANSWER": str(store)}):
                result = af.answer_to_doc("oc_1")
        self.assertIn("没有可转的回答", result)

    def test_creates_doc(self):
        payload = {"ok": True, "data": {"url": "https://x.feishu.cn/docx/abc"}}
        with tempfile.TemporaryDirectory() as tmp:
            store = Path(tmp) / "la.json"
            with patch.dict(os.environ, {"FEISHU_PARTNER_LAST_ANSWER": str(store)}):
                af.save_last_answer("oc_1", "周报总结", "本周完成 A、B、C 三项工作，下周继续推进。")
                with patch.object(af, "run_lark", return_value=payload) as mock_run:
                    result = af.answer_to_doc("oc_1", title="我的周报")
        self.assertIn("我的周报", result)
        self.assertIn("https://x.feishu.cn/docx/abc", result)
        argv = mock_run.call_args[0][0]
        self.assertEqual(argv[:2], ["docs", "+create"])
        self.assertIn("--doc-format", argv)

    def test_create_failure(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Path(tmp) / "la.json"
            with patch.dict(os.environ, {"FEISHU_PARTNER_LAST_ANSWER": str(store)}):
                af.save_last_answer("oc_1", "q", "一段足够长的回答内容，用来通过最小长度校验。")
                with patch.object(
                    af, "run_lark", return_value={"ok": False, "error": {"message": "denied"}}
                ):
                    result = af.answer_to_doc("oc_1")
        self.assertNotIn("已生成", result)


if __name__ == "__main__":
    unittest.main()
