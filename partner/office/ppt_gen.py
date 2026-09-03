"""PPT 生成技能：大纲 → pptx → 上传飞书。对标飞书豆包工作「PPT」内置技能。

大纲来源两种：Hermes LLM 根据主题+材料生成，或用户直接给 markdown 大纲文件。
pptx 用 python-pptx 本地生成；上传走 drive +import --type slides（转飞书在线幻灯片）。
安全默认：只生成不上传；--upload 显式开启才写飞书。
"""

from __future__ import annotations

import os
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from ..compose.formatters import format_lark_error
from ..core.lark import run_lark

CN_TZ = timezone(timedelta(hours=8))
PPT_DIR = Path.home() / ".feishu-partner" / "ppt"
MAX_BULLETS_PER_SLIDE = 6
MAX_SLIDES = 20


def ppt_dir() -> Path:
    override = os.environ.get("FEISHU_PARTNER_PPT")
    if override:
        return Path(override).expanduser()
    return PPT_DIR


def _slug(text: str) -> str:
    raw = re.sub(r"[^\w一-鿿-]+", "-", (text or "").strip())[:40].strip("-")
    return raw or "deck"


def parse_outline(text: str) -> dict[str, Any]:
    """把 markdown 式大纲解析成 {title, subtitle, slides:[{heading, bullets}]}。

    格式：
      # 演示标题
      副标题（可选，紧跟标题的一行普通文本）
      ## 页面标题
      - 要点
      - 要点
    """
    deck: dict[str, Any] = {"title": "", "subtitle": "", "slides": []}
    current: dict[str, Any] | None = None
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("## "):
            heading = line[3:].strip()
            if not heading:
                continue
            current = {"heading": heading, "bullets": []}
            deck["slides"].append(current)
        elif line.startswith("# "):
            if not deck["title"]:
                deck["title"] = line[2:].strip()
        elif line.startswith("- ") or line.startswith("* "):
            bullet = line[2:].strip()
            if not bullet:
                continue
            if current is not None and len(current["bullets"]) < MAX_BULLETS_PER_SLIDE:
                current["bullets"].append(bullet)
        else:
            # 标题后的第一行普通文本当副标题；页面里的普通行当要点
            if current is None and deck["title"] and not deck["subtitle"]:
                deck["subtitle"] = line
            elif current is not None and len(current["bullets"]) < MAX_BULLETS_PER_SLIDE:
                current["bullets"].append(line)
    deck["slides"] = [s for s in deck["slides"] if s["heading"]][:MAX_SLIDES]
    return deck


def outline_via_llm(topic: str, materials: str = "") -> str:
    """Hermes 生成大纲；禁用/失败返回空串走确定性模板。"""
    if os.environ.get("FEISHU_PARTNER_NO_LLM") == "1":
        return ""
    from ..compose.llm import _invoke_hermes

    prompt = (
        "你是飞书里的 PPT 助手。为主题写一份演示大纲，只输出大纲，不要解释。\n"
        "格式严格如下：\n"
        "# 演示标题\n"
        "一句话副标题\n"
        "## 页面一标题\n"
        "- 要点\n"
        "- 要点\n"
        "（每页 3-5 条要点，每条不超过 30 字；共 5-8 页；"
        "只根据材料，材料没有的信息不要编）\n\n"
        f"【主题】{topic}\n"
    )
    if materials.strip():
        prompt += f"\n【材料】\n{materials.strip()[:6000]}"
    raw = _invoke_hermes(prompt, timeout=90)
    text = (raw or "").strip()
    if not text or "Error:" in text[:80]:
        return ""
    # 至少要有一个 ## 页面才算可用
    if not re.search(r"^## ", text, re.M):
        return ""
    return text


def fallback_outline(topic: str) -> str:
    """无 LLM 时的确定性大纲骨架。"""
    t = (topic or "").strip() or "工作汇报"
    return (
        f"# {t}\n"
        "自动生成的大纲骨架，请按实际内容修改\n"
        "## 背景与目标\n- 现状概述\n- 本次目标\n"
        "## 核心进展\n- 进展一\n- 进展二\n"
        "## 问题与风险\n- 问题一\n- 应对策略\n"
        "## 下一步计划\n- 计划一\n- 计划二\n"
    )


def build_pptx(deck: dict[str, Any], *, filename: str = "") -> Path:
    """大纲 dict → 本地 .pptx。标题页 + 每页标题+要点。"""
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    title = (deck.get("title") or "").strip() or "未命名演示"
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    subtitle_shape = slide.placeholders[1] if len(slide.placeholders) > 1 else None
    if subtitle_shape is not None:
        subtitle_shape.text = (deck.get("subtitle") or "").strip()
    for spec in deck.get("slides") or []:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = spec["heading"]
        body = s.placeholders[1].text_frame
        body.clear()
        for index, bullet in enumerate(spec["bullets"]):
            para = body.paragraphs[0] if index == 0 else body.add_paragraph()
            para.text = bullet
            para.level = 0
            for run in para.runs:
                run.font.size = Pt(18)
    ppt_dir().mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(CN_TZ).strftime("%Y%m%d-%H%M%S")
    name = filename or f"{stamp}-{_slug(title)}.pptx"
    if not name.endswith(".pptx"):
        name += ".pptx"
    path = ppt_dir() / name
    prs.save(str(path))
    return path


def upload_slides(pptx_path: Path, *, name: str = "") -> str:
    """drive +import --type slides：本地 pptx → 飞书在线幻灯片。"""
    argv = [
        "drive",
        "+import",
        "--file",
        str(pptx_path),
        "--type",
        "slides",
    ]
    if name.strip():
        argv.extend(["--name", name.strip()])
    payload = run_lark(argv, as_identity="user", timeout=180)
    if payload.get("ok") is False or payload.get("error"):
        return "上传飞书失败：\n" + format_lark_error(payload)
    url = _pick_url(payload)
    if url:
        return f"已上传为飞书幻灯片：\n{url}"
    return "已上传，但没拿到链接（import 任务返回不完整）。"


def _pick_url(payload: dict[str, Any]) -> str:
    stack: list[Any] = [payload, payload.get("data")]
    while stack:
        cur = stack.pop(0)
        if not isinstance(cur, dict):
            continue
        for key in ("url", "doc_url", "link"):
            val = cur.get(key)
            if isinstance(val, str) and val.startswith("http"):
                return val
        for nested in ("data", "result", "task"):
            if isinstance(cur.get(nested), dict):
                stack.append(cur[nested])
    return ""


def generate_ppt(
    topic: str,
    *,
    outline_text: str = "",
    materials: str = "",
    upload: bool = False,
    name: str = "",
) -> str:
    """完整链路：大纲 → pptx →（可选）上传飞书。"""
    topic = (topic or "").strip()
    outline = outline_text.strip() or outline_via_llm(topic, materials) or fallback_outline(topic)
    deck = parse_outline(outline)
    if not deck["title"]:
        deck["title"] = topic or "未命名演示"
    if not deck["slides"]:
        return (
            "大纲里没有解析出任何页面（需要 ## 页面标题 + - 要点）。\n"
            "大纲格式：# 标题 / ## 页面标题 / - 要点"
        )
    outline_source = (
        "用户提供"
        if outline_text.strip()
        else ("LLM 生成" if outline != fallback_outline(topic) else "模板骨架")
    )
    path = build_pptx(deck)
    parts = [
        f"【{deck['title']}】共 {len(deck['slides'])} 页（大纲来源：{outline_source}）",
        f"已生成：{path}",
        "页目录：" + "、".join(s["heading"] for s in deck["slides"][:10]),
    ]
    if upload:
        parts.append(upload_slides(path, name=name or deck["title"]))
    else:
        parts.append("未上传。确认后用 --upload 转成飞书在线幻灯片。")
    return "\n\n".join(parts)


def ppt_cli(argv: list[str]) -> int:
    """CLI 入口：feishu ppt <主题> [options]"""
    import argparse

    parser = argparse.ArgumentParser(prog="feishu ppt", description="PPT 生成：大纲→pptx→上传飞书")
    parser.add_argument("topic", help="演示主题")
    parser.add_argument(
        "--outline", "-o", default="", help="大纲文本，或以 @ 开头的本地大纲文件路径"
    )
    parser.add_argument("--materials", "-m", default="", help="补充材料文本，或 @文件路径")
    parser.add_argument("--upload", action="store_true", help="上传到飞书（转在线幻灯片）")
    parser.add_argument("--name", default="", help="上传后的文档名")
    args = parser.parse_args(argv)

    def _resolve(text: str) -> str:
        if text.startswith("@"):
            p = Path(text[1:]).expanduser()
            if p.exists():
                return p.read_text(encoding="utf-8")
        return text

    print(
        generate_ppt(
            args.topic,
            outline_text=_resolve(args.outline),
            materials=_resolve(args.materials),
            upload=args.upload,
            name=args.name,
        )
    )
    return 0
