"""产物多样化：Excel 产物 + 回答一键转文档。对标飞书豆包工作「产物」能力。

- Excel：markdown 表格 → 本地 .xlsx（openpyxl，无则降级 CSV），--upload 转飞书在线表格。
- 转文档：serve 每轮回复后缓存「上一个回答」，用户说「转文档」即生成飞书云文档。
  缓存独立于 core/session（session 承载 follow-up 语义，不能被覆盖）。

安全默认：Excel 只生成不上传；转文档生成的是新文档，不改任何已有内容。
"""

from __future__ import annotations

import csv
import json
import os
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from ..compose.formatters import format_lark_error
from ..core.lark import run_lark
from .calendar_views import _created_doc_link

CN_TZ = timezone(timedelta(hours=8))
ARTIFACT_DIR = Path.home() / ".feishu-partner" / "artifacts"
_ANSWER_STORE = Path.home() / ".feishu-partner" / "last-answer.json"
_ANSWER_TTL = timedelta(hours=2)
MIN_ANSWER_LEN = 10

# serve 里的触发词：用户接着说这些 → 把上一个回答转成云文档
DOC_TRIGGERS = ("转文档", "转成文档", "生成文档", "存成文档", "写成文档", "导出文档", "存为文档")


def artifact_dir() -> Path:
    override = os.environ.get("FEISHU_PARTNER_ARTIFACT_DIR")
    if override:
        return Path(override).expanduser()
    return ARTIFACT_DIR


def _answer_store_path() -> Path:
    override = os.environ.get("FEISHU_PARTNER_LAST_ANSWER")
    if override:
        return Path(override).expanduser()
    return _ANSWER_STORE


def _slug(text: str) -> str:
    raw = re.sub(r"[^\w一-鿿-]+", "-", (text or "").strip())[:40].strip("-")
    return raw or "artifact"


# ---------------------------------------------------------------- markdown 表格解析

_TABLE_ROW_RE = re.compile(r"^\s*\|(.+)\|\s*$")
_TABLE_SEP_RE = re.compile(r"^\s*\|[\s:|-]+\|\s*$")


def parse_markdown_tables(text: str) -> list[list[list[str]]]:
    """从文本里抽出所有 markdown 表格 → [表][行][单元格]。首行为表头。"""
    tables: list[list[list[str]]] = []
    current: list[list[str]] = []
    for line in (text or "").splitlines():
        if _TABLE_SEP_RE.match(line):
            continue  # |---|---| 分隔行
        match = _TABLE_ROW_RE.match(line)
        if match:
            cells = [c.strip() for c in match.group(1).split("|")]
            current.append(cells)
            continue
        if current:
            tables.append(current)
            current = []
    if current:
        tables.append(current)
    # 至少要表头 + 一行数据，且表头非全空
    return [t for t in tables if len(t) >= 2 and any(t[0])]


_NUM_RE = re.compile(r"^-?[\d,]+(\.\d+)?%?$")


def _maybe_number(cell: str) -> Any:
    raw = (cell or "").strip()
    if _NUM_RE.match(raw):
        is_pct = raw.endswith("%")
        body = raw.rstrip("%").replace(",", "")
        try:
            value: Any = float(body) if "." in body else int(body)
        except ValueError:
            return cell
        if is_pct:
            return value / 100.0
        return value
    return cell


# ---------------------------------------------------------------- Excel 产物


def build_xlsx(tables: list[list[list[str]]], *, filename: str = "") -> Path:
    """markdown 表格组 → 本地 .xlsx。多表 → 多 sheet。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    for index, table in enumerate(tables):
        sheet = wb.active if index == 0 else wb.create_sheet()
        sheet.title = f"表{index + 1}" if len(tables) > 1 else "数据"
        widths: dict[int, int] = {}
        for row_idx, row in enumerate(table):
            for col_idx, cell in enumerate(row):
                target = sheet.cell(row=row_idx + 1, column=col_idx + 1)
                value = _maybe_number(cell) if row_idx > 0 else cell
                target.value = value
                if row_idx == 0:
                    target.font = Font(bold=True)
                widths[col_idx] = max(widths.get(col_idx, 0), len(str(cell)) + 2)
        for col_idx, width in widths.items():
            sheet.column_dimensions[get_column_letter(col_idx + 1)].width = min(
                max(width, 8), 40
            )
    artifact_dir().mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(CN_TZ).strftime("%Y%m%d-%H%M%S")
    name = filename or f"{stamp}-table.xlsx"
    if not name.endswith(".xlsx"):
        name += ".xlsx"
    path = artifact_dir() / name
    wb.save(str(path))
    return path


def build_csv_fallback(table: list[list[str]], *, filename: str = "") -> Path:
    """无 openpyxl 时的降级：第一张表写 CSV（Excel/WPS 可直接打开）。"""
    artifact_dir().mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(CN_TZ).strftime("%Y%m%d-%H%M%S")
    name = filename or f"{stamp}-table.csv"
    if not name.endswith(".csv"):
        name += ".csv"
    path = artifact_dir() / name
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        csv.writer(fh).writerows(table)
    return path


def upload_sheet(xlsx_path: Path, *, name: str = "") -> str:
    """drive +import --type sheet：本地 xlsx → 飞书在线表格。"""
    argv = ["drive", "+import", "--file", str(xlsx_path), "--type", "sheet"]
    if name.strip():
        argv.extend(["--name", name.strip()])
    payload = run_lark(argv, as_identity="user", timeout=180)
    if payload.get("ok") is False or payload.get("error"):
        return "上传飞书失败：\n" + format_lark_error(payload)
    url = _pick_url(payload)
    if url:
        return f"已上传为飞书表格：\n{url}"
    return "已上传，但没拿到链接（import 任务返回不完整）。"


def _pick_url(payload: dict[str, Any]) -> str:
    stack: list[Any] = [payload, payload.get("data")]
    while stack:
        cur = stack.pop(0)
        if not isinstance(cur, dict):
            continue
        for key in ("url", "doc_url", "link"):
            val = cur.get(key)
            if isinstance(val, str) and val.startswith("http"):
                return val
        for nested in ("data", "result", "task"):
            if isinstance(cur.get(nested), dict):
                stack.append(cur[nested])
    return ""


def make_excel(source: str, *, upload: bool = False, name: str = "") -> str:
    """完整链路：markdown 表格文本（或 @文件）→ xlsx →（可选）上传飞书。"""
    text = _resolve_source(source)
    tables = parse_markdown_tables(text)
    if not tables:
        return (
            "没解析出 markdown 表格。格式：\n"
            "| 列1 | 列2 |\n| --- | --- |\n| a | 1 |"
        )
    try:
        path = build_xlsx(tables, filename=_slug(name) if name else "")
        kind = "xlsx"
    except ImportError:
        path = build_csv_fallback(tables[0], filename=_slug(name) if name else "")
        kind = "csv"
    parts = [
        f"共解析 {len(tables)} 张表，最大 {max(len(t) for t in tables) - 1} 行数据",
        f"已生成：{path}",
    ]
    if upload and kind == "xlsx":
        parts.append(upload_sheet(path, name=name))
    elif upload:
        parts.append("CSV 降级产物不支持转在线表格，未上传。")
    else:
        parts.append("未上传。确认后用 --upload 转成飞书在线表格。")
    return "\n\n".join(parts)


def _resolve_source(source: str) -> str:
    raw = (source or "").strip()
    if raw.startswith("@"):
        path = Path(raw[1:]).expanduser()
        if path.exists():
            return path.read_text(encoding="utf-8")
        return ""
    return raw


# ---------------------------------------------------------------- 回答一键转文档


def save_last_answer(chat_id: str, query: str, answer: str) -> None:
    """serve 每轮回复成功后调用。独立存储，不碰 session 的 follow-up 语义。"""
    cid = (chat_id or "").strip()
    text = (answer or "").strip()
    if not cid or len(text) < MIN_ANSWER_LEN:
        return
    path = _answer_store_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    blob: dict[str, Any] = {}
    if path.exists():
        try:
            loaded = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            loaded = {}
        if isinstance(loaded, dict):
            blob = loaded
    blob[cid] = {
        "ts": datetime.now(CN_TZ).isoformat(timespec="seconds"),
        "query": (query or "")[:500],
        "answer": text[:20000],
    }
    path.write_text(json.dumps(blob, ensure_ascii=False, indent=2), encoding="utf-8")


def load_last_answer(chat_id: str) -> dict[str, str] | None:
    cid = (chat_id or "").strip()
    if not cid:
        return None
    path = _answer_store_path()
    if not path.exists():
        return None
    try:
        blob = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    row = blob.get(cid) if isinstance(blob, dict) else None
    if not isinstance(row, dict):
        return None
    try:
        when = datetime.fromisoformat(str(row.get("ts") or ""))
    except ValueError:
        return None
    if when.tzinfo is None:
        when = when.replace(tzinfo=CN_TZ)
    if datetime.now(CN_TZ) - when.astimezone(CN_TZ) > _ANSWER_TTL:
        return None
    if not str(row.get("answer") or "").strip():
        return None
    return row


def looks_like_doc_request(raw: str) -> bool:
    text = re.sub(r"[？?。！!，,\s]+", "", (raw or "").strip())
    if not text or len(text) > 10:
        return False
    return any(word in text for word in DOC_TRIGGERS)


def answer_to_doc(chat_id: str, *, title: str = "") -> str:
    """把上一个回答转成飞书云文档，返回链接。"""
    row = load_last_answer(chat_id)
    if not row:
        return "最近 2 小时内没有可转的回答。先问点什么，再说「转文档」。"
    answer = str(row["answer"]).strip()
    query = str(row.get("query") or "").strip()
    doc_title = title.strip() or (query[:30] if query else "对话记录")
    doc_title = f"{doc_title} · {datetime.now(CN_TZ).date().isoformat()}"
    header = f"> 来源：{query}\n\n" if query else ""
    created = run_lark(
        [
            "docs",
            "+create",
            "--title",
            doc_title,
            "--doc-format",
            "markdown",
            "--content",
            header + answer,
        ],
        as_identity="user",
    )
    if created.get("ok"):
        link = _created_doc_link(created)
        extra = f"\n{link}" if link else ""
        return f"已生成云文档《{doc_title}》。{extra}\n聊天里不贴正文，打开文档看。"
    return format_lark_error(created)


# ---------------------------------------------------------------- CLI


def artifact_cli(argv: list[str]) -> int:
    """CLI 入口：feishu artifact <excel|doc> ..."""
    import argparse

    parser = argparse.ArgumentParser(
        prog="feishu artifact", description="产物多样化：Excel 产物 / 回答转文档"
    )
    sub = parser.add_subparsers(dest="kind", required=True)

    p_excel = sub.add_parser("excel", help="markdown 表格 → xlsx（可选上传飞书）")
    p_excel.add_argument("source", help="markdown 表格文本，或 @文件路径")
    p_excel.add_argument("--upload", action="store_true", help="上传到飞书（转在线表格）")
    p_excel.add_argument("--name", default="", help="文件名 / 上传后的文档名")

    p_doc = sub.add_parser("doc", help="把缓存的上一个回答转成云文档")
    p_doc.add_argument("--title", default="", help="文档标题（默认取原问题）")
    p_doc.add_argument("--chat", default="cli", help="chat_id（默认 cli）")

    args = parser.parse_args(argv)
    if args.kind == "excel":
        print(make_excel(args.source, upload=args.upload, name=args.name))
        return 0
    print(answer_to_doc(args.chat, title=args.title))
    return 0
